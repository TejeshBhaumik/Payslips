#library to import the excel file
import openpyxl
import shutil

import sendEmail



#libraries to create the pdf file and add text to it
from reportlab.pdfgen import canvas

from reportlab.lib.styles import (ParagraphStyle, getSampleStyleSheet)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfbase.ttfonts import TTFont
#libraries to merge pdf files

from reportlab.lib.units import inch, cm
import os
from PyPDF2 import PdfReader, PdfWriter

from reportlab.lib import pdfencrypt

from reportlab.platypus import Table, TableStyle, Paragraph

from reportlab.platypus import SimpleDocTemplate, Spacer, Image
from reportlab.lib.pagesizes import A4, A5, letter

from reportlab.platypus import TableStyle
from reportlab.lib import colors

from flask import Flask, render_template, request, jsonify, send_file
from zipfile import ZipFile
from io import BytesIO
import os
import requests
import logging
import threading
import uuid
import json
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

app = Flask(__name__)
logging.basicConfig(level=logging.INFO)
logging.getLogger("werkzeug").setLevel(logging.WARNING)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SMTP_RETRY_DELAY_SECONDS = int(os.environ.get("SMTP_RETRY_DELAY_SECONDS", "600"))
JOB_STATE_DIR = os.environ.get("JOB_STATE_DIR", os.path.join(tempfile.gettempdir(), "payslips_jobs"))
JOBS = {}
JOBS_LOCK = threading.Lock()
JOB_EXECUTOR = ThreadPoolExecutor(max_workers=int(os.environ.get("JOB_WORKER_THREADS", "2")))


def job_state_path(job_id):
    return os.path.join(JOB_STATE_DIR, f"{job_id}.json")


def write_job_state(job):
    os.makedirs(JOB_STATE_DIR, exist_ok=True)
    path = job_state_path(job["id"])
    tmp_path = f"{path}.{uuid.uuid4().hex}.tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(job, f)
    os.replace(tmp_path, path)
    app.logger.info(
        "Job %s persisted status=%s sent=%s total=%s current=%s",
        job.get("id"),
        job.get("status"),
        job.get("sent"),
        job.get("total"),
        job.get("current"),
    )


def read_job_state(job_id):
    try:
        with open(job_state_path(job_id), "r", encoding="utf-8") as f:
            job = json.load(f)
            app.logger.info("Job %s loaded from disk status=%s sent=%s", job_id, job.get("status"), job.get("sent"))
            return job
    except FileNotFoundError:
        app.logger.warning("Job %s not found in memory or disk", job_id)
        return None


def create_job():
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = {
            "id": job_id,
            "status": "queued",
            "sent": 0,
            "total": 0,
            "current": "",
            "message": "Preparing batch.",
            "error": "",
            "resume_row": None,
            "retry_at": "",
            "retry_count": 0,
        }
        write_job_state(JOBS[job_id])
        app.logger.info("Job %s created", job_id)
    return job_id


def update_job(job_id, **changes):
    if not job_id:
        return
    with JOBS_LOCK:
        if job_id in JOBS:
            JOBS[job_id].update(changes)
            write_job_state(JOBS[job_id])
        else:
            app.logger.warning("Job %s update skipped; job not found. changes=%s", job_id, sorted(changes.keys()))


def get_job(job_id):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    return job or read_job_state(job_id)


def count_rows_to_send(sheet, start_row=3):
    count = 0
    for row_number in range(start_row, sheet.max_row + 1):
        if sheet.cell(row=row_number, column=1).value is None:
            break
        count += 1
    return count


def schedule_retry(
    job_id,
    form_data,
    excel_bytes,
    excel_filename,
    excel_content_type,
    image_filename,
    resume_row,
    sent,
):
    retry_at = datetime.now(timezone.utc) + timedelta(seconds=SMTP_RETRY_DELAY_SECONDS)
    job = get_job(job_id) or {}
    retry_count = int(job.get("retry_count") or 0) + 1
    retry_delay_label = (
        f"{SMTP_RETRY_DELAY_SECONDS // 60} minutes"
        if SMTP_RETRY_DELAY_SECONDS >= 60
        else f"{SMTP_RETRY_DELAY_SECONDS} seconds"
    )
    message = (
        f"Email provider stopped sending. Retrying from Excel row {resume_row} "
        f"in {retry_delay_label}."
    )
    update_job(
        job_id,
        status="waiting",
        sent=sent,
        resume_row=resume_row,
        retry_at=retry_at.isoformat(),
        retry_count=retry_count,
        message=message,
        error="",
        current=f"Excel row {resume_row}",
    )
    app.logger.warning(
        "Batch %s paused at Excel row %s after sent=%s; retrying at %s",
        job_id,
        resume_row,
        sent,
        retry_at.isoformat(),
    )

    timer = threading.Timer(
        SMTP_RETRY_DELAY_SECONDS,
        lambda: JOB_EXECUTOR.submit(
            run_extract_job,
            job_id,
            form_data,
            excel_bytes,
            excel_filename,
            excel_content_type,
            image_filename,
            resume_row,
            sent,
        ),
    )
    timer.daemon = True
    timer.start()

@app.route('/')
def index():
  return render_template("index.html")

@app.route('/healthz')
def healthz():
    return "ok", 200

@app.route('/favicon.ico')
def favicon():
    return "", 204

@app.route('/extract', methods=["POST"])
def create_payslip():
    form_data = request.form.to_dict()
    excel_file = request.files.get('excelFile')
    image_file = request.files.get('image')

    if excel_file is None or excel_file.filename == "":
        app.logger.warning("No employee data file uploaded.")
        return jsonify({"error": "please enter employee data excel"}), 400

    job_id = create_job()
    excel_bytes = excel_file.read()
    excel_filename = excel_file.filename
    excel_content_type = excel_file.content_type
    image_filename = image_file.filename if image_file and image_file.filename else ""

    app.logger.info(
        "POST /extract accepted job=%s option=%s filename=%s bytes=%d",
        job_id,
        form_data.get("options"),
        excel_filename,
        len(excel_bytes),
    )
    future = JOB_EXECUTOR.submit(run_extract_job, job_id, form_data, excel_bytes, excel_filename, excel_content_type, image_filename)
    app.logger.info("Job %s submitted to executor future=%s", job_id, id(future))

    return jsonify({"jobId": job_id}), 202


@app.route('/extract/status/<job_id>')
def extract_status(job_id):
    job = get_job(job_id)
    if job is None:
        return jsonify({"error": "Batch not found."}), 404
    app.logger.info(
        "GET /extract/status/%s status=%s sent=%s total=%s current=%s",
        job_id,
        job.get("status"),
        job.get("sent"),
        job.get("total"),
        job.get("current"),
    )
    return jsonify(job)


def run_extract_job(
    job_id,
    form_data,
    excel_bytes,
    excel_filename,
    excel_content_type,
    image_filename,
    start_row=None,
    sent_so_far=0,
):
    app.logger.info("Job %s worker started start_row=%s sent_so_far=%s", job_id, start_row, sent_so_far)
    try:
        message, status_code = run_extract_sync(
            form_data,
            excel_bytes,
            excel_filename,
            excel_content_type,
            image_filename,
            job_id,
            start_row=start_row,
            sent_so_far=sent_so_far,
        )
        if status_code >= 400:
            update_job(job_id, status="failed", error=message, message=message)
            return
        update_job(
            job_id,
            status="complete",
            message="all payslips generated and sent",
            current="",
        )
        app.logger.info("Batch %s completed.", job_id)
    except sendEmail.EmailRetryNeeded as e:
        resume_row = getattr(e, "excel_row", start_row)
        sent = getattr(e, "sent", sent_so_far)
        if resume_row is None:
            message = str(e) or "Email provider stopped sending. Please retry later."
            update_job(job_id, status="limited", error=message, message=message, current="")
            return
        schedule_retry(
            job_id,
            form_data,
            excel_bytes,
            excel_filename,
            excel_content_type,
            image_filename,
            resume_row,
            sent,
        )
    except Exception as e:
        app.logger.exception("Batch %s failed", job_id)
        update_job(job_id, status="failed", error=str(e), message=str(e), current="")


def run_extract_sync(
    form_data,
    excel_bytes,
    excel_filename,
    excel_content_type,
    image_filename,
    job_id,
    start_row=None,
    sent_so_far=0,
):

    #this was changed
    #convert the font so it is compatible
    pdfmetrics.registerFont(TTFont('Arial', os.path.join(BASE_DIR, 'arial.ttf')))
    file_header = ""
    bottom_label = ""
    email_file_body = ""
    is_encrypted = form_data.get('encrypt')

    option = form_data.get('options')
    smtp_email = form_data.get('smtpEmail', "")
    smtp_password = form_data.get('smtpPassword', "")
    if option not in {"custom", "salarySlip", "emailOnly"}:
        app.logger.warning("Invalid workflow option=%s", option)
        return "Please choose a workflow before starting the batch.", 400
    
    # Handle the uploaded file and form inputs here
    # Example: Save the file, process it, etc.

    if not excel_bytes:
        app.logger.warning("No employee data file uploaded.")
        return "please enter employee data excel", 400
    app.logger.info(
        "Batch %s started. option=%s file=%s bytes=%d start_row=%s sent_so_far=%s",
        job_id,
        option,
        excel_filename,
        len(excel_bytes),
        start_row,
        sent_so_far,
    )
    

    if option == 'custom':
        file_header = form_data['field1']
        bottom_label = form_data['field2']
        email_file_body = form_data['field3']
        # Handle custom fields data
    
    if option == "emailOnly":

        img = image_filename
        sender = form_data.get('field11', "")
        sender_title = form_data.get('field12', "")
        company = form_data.get('field13', "")
        com_address = form_data.get('field14', "")
        ph_number = form_data.get('field15', "")
        com_email = form_data.get('field16', "")
    
    if option == 'salarySlip':
        file_header = "Salary slip for the month of"
        email_file_body = "salary slip for the month of"
        bottom_label = "The figures in the salary slip are confidential and not to be disclosed.\
              Signature is not required for this payslip"

    #import the sheet from the excel file
    try:
        wb = openpyxl.load_workbook(BytesIO(excel_bytes), data_only=True)
        sheet = wb[wb.sheetnames[0]]
    except Exception as e:
        app.logger.exception("Error loading workbook")
        return f"Error loading workbook: {e}", 400

    if option == "emailOnly":
        first_row = start_row or 1
        email_send_total = count_rows_to_send(sheet, start_row=1)
        app.logger.info("Batch %s loaded workbook. rows=%d workflow=emailOnly", job_id, email_send_total)
        email_send_index = sent_so_far
        update_job(
            job_id,
            status="running",
            total=email_send_total,
            sent=email_send_index,
            message="Sending email batch.",
            retry_at="",
            resume_row=None,
        )
        for row_number in range(first_row, sheet.max_row + 1):
            name = sheet.cell(row=row_number, column=1)
            if name.value is None:
                break
            email = sheet.cell(row=row_number, column=2)
            app.logger.info(
                "Job %s emailOnly row=%s send_index=%s recipient=%s",
                job_id,
                row_number,
                email_send_index + 1,
                email.value,
            )
            try:
                email_send_index += 1
                update_job(
                    job_id,
                    current=str(email.value),
                    message=f"Sending {email_send_index} of {email_send_total}",
                )
                sendEmail.sendEmailWithImage(
                    str(name.value),
                    img,
                    sender,
                    sender_title,
                    company,
                    com_address,
                    ph_number,
                    com_email,
                    str(email.value),
                    smtp_user=smtp_email,
                    smtp_password=smtp_password,
                )
                update_job(
                    job_id,
                    sent=email_send_index,
                    current=str(email.value),
                    message=f"Sent {email_send_index} of {email_send_total}",
                )

            except sendEmail.EmailRetryNeeded as e:
                email_send_index -= 1
                app.logger.warning(
                    "Job %s emailOnly paused at row=%s sent=%s reason=%s",
                    job_id,
                    row_number,
                    email_send_index,
                    e,
                )
                e.excel_row = row_number
                e.sent = email_send_index
                raise
            except Exception as e:
                app.logger.exception("emailOnly send failed at row=%s", row_number)
                return f"Email send failed at row {row_number}: {e}", 400
        return "all payslips generated and sent", 200






    #Page information
    page_width = 2156
    page_height = 3050
    spread = 50
    start = 200
    start_2 = 700

    #Payslip variables
    company_name = 'Conacent Consulting'

    logo = Image(os.path.join(BASE_DIR, "logo and address.png"))
    logo._restrictSize(6*inch,8*inch)
    logo.hAlign = "LEFT"
    logo.VALIGN = "TOP"
    col_names = []
    for name in sheet[2]:
        if name.value is not None and name.value != "email":
            col_names.append(name.value)
        

    super_columns = []
    for super_column in sheet[1]:
        super_columns.append(super_column.value)

    year = ""

    styles = getSampleStyleSheet()
    i = start_row or 3
    pdf_send_total = count_rows_to_send(sheet, start_row=3)
    app.logger.info("Batch %s loaded workbook. rows=%d workflow=%s", job_id, pdf_send_total, option)
    pdf_send_index = sent_so_far
    update_job(
        job_id,
        status="running",
        total=pdf_send_total,
        sent=pdf_send_index,
        message="Generating and sending payslips.",
        retry_at="",
        resume_row=None,
    )
    while (i):
        vals = []
        email = ""
        password = ""
        if sheet.cell(row = i, column = 1).value is None:
            break
        row_number = i
        for j in range(1,sheet.max_column+1):
           
            if sheet.cell(row = i, column = j).value is not None:

                

                inp = str(sheet.cell(row = i, column = j).value)
        
                if "00:00:00" in inp:

                    inp = inp.replace("00:00:00","")    
                      
                    y = inp[:inp.find("-")]
                    # inp = inp[::-1]
                    month = inp[inp.find("-") + 1: inp.rfind("-")]
                    date = inp[inp.rfind("-") + 1: ]
                    inp = date + "/" + month + "/" + y
                    inp = inp.replace(" ","")             
                if "@" in inp:
                    email = inp
                else :
                    vals.append(inp)
            else:
                vals.append("N/A")
        name =  str(vals[1])+ ' ' + str(vals[0])  + '.pdf' 
        pdf_buffer = BytesIO()
        pdf = SimpleDocTemplate(
                    pdf_buffer,
                    pagesize= A4,
                    )

        if str(vals[2]) != 'x' and str(vals[2]) != 'N/A' and type(vals[0]) == str:
            app.logger.info("Job %s building PDF row=%s name=%s recipient=%s", job_id, row_number, name, email)
            data = []
            elements = [logo]
            for j in range(len(vals)):
                if super_columns[j] is None:
                    if str(vals[j]) == 'N/A':
                        continue
                    data.append([col_names[j], str(vals[j])])
                elif super_columns[j] == "LEAVE STATEMENT":
                    table = Table(data, rowHeights = len(data) * [13], colWidths=inch*3)
                    style = TableStyle([
                            ('BACKGROUND', (0,0), (3,0), colors.steelblue),
                            ('TEXTCOLOR',(0,0),(-1,0),colors.whitesmoke),

                            ('ALIGN',(0,0),(-1,-1),'CENTER'),

                            ('FONTNAME', (0,0), (-1,0), 'Courier-Bold'),
                            ('FONTSIZE', (0,0), (-1,0), 10),

                            ('BOTTOMPADDING', (0,0), (-1,0), 2),

                            ('BACKGROUND',(0,1),(-1,-1),colors.beige),
                        ])
                    table.hAlign = "LEFT"
                    
                    table.setStyle(style)


                    elements.append(table)
                    data = []
                    data.append([super_columns[j], ""])
                    data.append([col_names[j], str(vals[j])])
                else:
                    if len(data) == 1:
                        continue
                    #This changes main table
                    #13 IS BEST FONT SIZE
                    table = Table(data, rowHeights = len(data) * [13], colWidths=inch*3)
                    if data[0][1] != "":

                        #First part
                        
                        data[0] = [Paragraph(f"<b> {file_header} " + " " + data[0][0].upper() +  "</b>")]
                        table = Table(data)
                        style = TableStyle([
                        ('BACKGROUND', (0,0), (3,0), colors.white),
                        ('TEXTCOLOR',(0,0),(-1,0),colors.black),

                        ('ALIGN',(0,0),(-1,-1),'LEFT'),

                        ('VALIGN',(-200,-200),(-100,-100),'TOP'),

                        ('FONTNAME', (0,0), (-1,0), 'Courier-Bold'),
                        ('FONTSIZE', (0,0), (-1,0), 100),
                        ('TEXTFONT', (0, 1), (-1, 1), 'Times-Bold'),

                        ('BOTTOMPADDING', (0,0), (-1,0), 2),

                        ('BACKGROUND',(0,1),(-1,-1),colors.white),
                        ])
                        table.hAlign = "LEFT"
                    else:         
                        
                        style = TableStyle([
                            ('BACKGROUND', (0,0), (3,0), colors.steelblue),
                            ('TEXTCOLOR',(0,0),(-1,0),colors.whitesmoke),

                            ('ALIGN',(0,0),(-1,-1),'CENTER'),

                            ('FONTNAME', (0,0), (-1,0), 'Courier-Bold'),
                            #('FONTSIZE', (0,0), (-1,0), 12),
                            ('TEXTSIZE', (0,0), (-1,0), 10),
                            #('TOPPADDING', (0,0), (-1,0), 2),
                            ('BOTTOMPADDING', (0,0), (-1,0), 2),

                            ('BACKGROUND',(0,1),(-1,-1),colors.beige),
                        ])
                        table.hAlign = "LEFT"
                    table.setStyle(style)
                    elements.append(table)
                    elements.append(Spacer(1,5))
                    data = []
                    data.append([super_columns[j], ""])
                    elements.append(Spacer(1,10))
                    data.append([col_names[j], vals[j]])
            
            #this is for leave 
            elements.append(Spacer(1,5))
            table = Table(data, rowHeights = len(data) * [13], colWidths=inch*3)
            style = TableStyle([
                                        ('BACKGROUND', (0,0), (3,0), colors.steelblue),
                                        ('TEXTCOLOR',(0,0),(-1,0),colors.whitesmoke),

                                        ('ALIGN',(0,0),(-1,-1),'CENTER'),

                                        ('FONTNAME', (0,0), (-1,0), 'Courier-Bold'),
                                        ('BOTTOMPADDING', (0,0), (-1,0), 2),

                                        ('BACKGROUND',(0,1),(-1,-1),colors.beige),
                                    ])
            table.hAlign = "LEFT"
            table.setStyle(style)
            elements.append(Spacer(1,10))
            elements.append(table)
            elements.append(Spacer(1,100))
            style_new = getSampleStyleSheet()
            yourStyle = ParagraphStyle('yourtitle',
                           fontName="Helvetica",
                           fontSize=8,
                           parent=style_new['Heading2'],
                           alignment=1,
                           spaceAfter=2)
            elements.append(Paragraph(f"<i>{bottom_label}</i>",  yourStyle))
            elements.append(Paragraph("<i> Registered Office:  P-94/95, Bangur Avenue, BL-C, Kolkata - 700055 </i>",  yourStyle))
            pdf.build(elements)
            pdf_buffer.seek(0)

            out = PdfWriter()

            # Read the generated PDF from memory
            filename = PdfReader(pdf_buffer)
            out.append_pages_from_reader(filename)
            is_pan = ""
            if is_encrypted == "on":
                password = vals[5]
                out.encrypt(password)
                is_pan = "Please use your PAN number as the password for opening the pdf document."

            out_buffer = BytesIO()
            out.write(out_buffer)
            out_buffer.seek(0)
            pdf_bytes = out_buffer.getvalue()
            pdf_buffer.close()
            out_buffer.close()


            year = vals[0]
            if (year != "N/A"):
                pdf_send_index += 1
                update_job(
                    job_id,
                    current=name,
                    message=f"Connecting to email server for payslip {pdf_send_index} of {pdf_send_total}",
                )
                app.logger.info(
                    "Job %s sending PDF row=%s send_index=%s/%s pdf=%s recipient=%s",
                    job_id,
                    row_number,
                    pdf_send_index,
                    pdf_send_total,
                    name,
                    email,
                )
                try:
                    sendEmail.sendEmailWithPDF(
                        pdf_bytes=pdf_bytes,
                        pdf_name=name,
                        email=email,
                        month=str(vals[0]),
                        person_name=str(vals[1]),
                        email_file_body=email_file_body,
                        is_pan=is_pan,
                        smtp_user=smtp_email,
                        smtp_password=smtp_password,
                    )
                    update_job(
                        job_id,
                        sent=pdf_send_index,
                        current=name,
                        message=f"Sent payslip {pdf_send_index} of {pdf_send_total}",
                    )
                except sendEmail.EmailRetryNeeded as e:
                    pdf_send_index -= 1
                    app.logger.warning(
                        "Job %s paused at row=%s sent=%s reason=%s",
                        job_id,
                        row_number,
                        pdf_send_index,
                        e,
                    )
                    e.excel_row = row_number
                    e.sent = pdf_send_index
                    raise
                except Exception as e:
                    app.logger.exception("Failed sending PDF email for %s", name)
                    return f"Email send failed for {name}: {e}", 400
        i += 1

          
    return "all payslips generated and sent", 200
        
if __name__ == "__main__":
    app.run(debug = True, threaded=True, port = int(os.environ.get('PORT', 5000)))
